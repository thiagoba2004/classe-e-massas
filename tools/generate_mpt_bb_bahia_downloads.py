from pathlib import Path
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "mpt" / "modelo-bb-bahia-greve-2026.html"
OUTDIR = ROOT / "mpt" / "downloads"
DOCX = OUTDIR / "Modelo_Denuncia_MPT_BB_Bahia_Greve_2026.docx"


def set_font(run, size=11, bold=None, italic=None):
    name = "Liberation Sans"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, node):
    if isinstance(node, NavigableString):
        if str(node):
            r = paragraph.add_run(str(node))
            set_font(r)
        return
    if not isinstance(node, Tag):
        return
    if node.name == "br":
        paragraph.add_run().add_break()
        return
    text = node.get_text()
    r = paragraph.add_run(text)
    set_font(r, bold=node.name in {"strong", "b"}, italic=node.name in {"em", "i"})


def add_html_paragraph(doc, tag, heading=False, centered=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(7)
    if not heading:
        p.paragraph_format.first_line_indent = Cm(0.7)
    else:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_with_next = True
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for child in tag.children:
        add_inline(p, child)
    for r in p.runs:
        if heading:
            set_font(r, size=13, bold=True)
    return p


def add_callout(doc, tag):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    for child in tag.children:
        if isinstance(child, Tag) and child.name == "p":
            for grand in child.children:
                add_inline(p, grand)
        else:
            add_inline(p, child)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    OUTDIR.mkdir(parents=True, exist_ok=True)
    soup = BeautifulSoup(HTML.read_text(encoding="utf-8"), "html.parser")
    petition = soup.select_one(".petition")
    if petition is None:
        raise SystemExit("Bloco .petition não encontrado no HTML")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Liberation Sans")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Liberation Sans")
    normal.font.size = Pt(11)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("CLASSE E MASSAS  |  MODELO DE DENÚNCIA AO MPT")
    set_font(r, size=8, bold=True)
    r.font.color.rgb = RGBColor(90, 90, 90)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Classe e Massas • 2026 • Modelo editável - confira os fatos e documentos antes do protocolo")
    set_font(r, size=8)
    r.font.color.rgb = RGBColor(95, 95, 95)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BANCO DO BRASIL • BAHIA • CAMPANHA SALARIAL 2026")
    set_font(r, size=9, bold=True)
    r.font.color.rgb = RGBColor(90, 90, 90)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MODELO DE DENÚNCIA AO MPT")
    set_font(r, size=22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("Greve, possível retaliação econômica, tratamento diferenciado entre bases e formação livre e informada da vontade coletiva")
    set_font(r, size=12, bold=True)

    warning = soup.select_one(".model-warning")
    if warning:
        add_callout(doc, warning)

    for child in petition.children:
        if not isinstance(child, Tag):
            continue
        classes = child.get("class", [])
        if child.name == "h2":
            centered = "petition-title" in classes
            add_html_paragraph(doc, child, heading=True, centered=centered)
        elif child.name == "p":
            centered = "petition-subtitle" in classes
            p = add_html_paragraph(doc, child, centered=centered)
            if "petition-address" in classes:
                p.paragraph_format.first_line_indent = Cm(0)
                for r in p.runs:
                    r.bold = True
        elif child.name == "div" and "model-box" in classes:
            add_callout(doc, child)
        elif child.name in {"ul", "ol"}:
            for li in child.find_all("li", recursive=False):
                p = doc.add_paragraph(style="List Bullet" if child.name == "ul" else "List Number")
                p.paragraph_format.space_after = Pt(4)
                for n in li.children:
                    add_inline(p, n)

    props = doc.core_properties
    props.title = "Modelo de Denúncia ao MPT - BB Bahia e Greve 2026"
    props.subject = "Possíveis práticas antissindicais e retaliação econômica em contexto de greve"
    props.author = "Classe e Massas"
    props.comments = "Modelo editável. Conferir fatos e anexos antes do protocolo."
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
